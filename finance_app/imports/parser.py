"""Strict CSV/XLSX portfolio parsing without any persistence side effects."""

from __future__ import annotations

import csv
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook  # type: ignore[import-untyped]

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 2_000
MAX_IMPORT_COLUMNS = 64
MAX_IMPORT_CELLS = 50_000
MAX_IMPORT_FIELD_CHARACTERS = 10_000
MAX_PREVIEW_CHARACTERS = 2 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_XLSX_ENTRY_BYTES = 25 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 200
_MIN_RATIO_CHECK_BYTES = 1_024
ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".png", ".jpg", ".jpeg"}

_ALIASES = {
    "基金代码": "asset_code",
    "代码": "asset_code",
    "基金名称": "asset_name",
    "名称": "asset_name",
    "份额": "quantity",
    "成本金额": "cost",
    "持仓成本": "cost",
    "当前市值": "market_value",
    "持有金额": "market_value",
    "可用现金": "available_cash",
    "日期": "date",
}
_MONEY = re.compile(r"(?:0|[1-9][0-9]*)(?:\.[0-9]{1,2})?")
_QUANTITY = re.compile(r"(?:0|[1-9][0-9]*)(?:\.[0-9]{1,8})?")
_CODE = re.compile(r"[A-Za-z0-9][A-Za-z0-9._-]{0,63}")
MAX_CENTS = 2**63 - 1
_MAX_CENTS_TEXT = str(MAX_CENTS)
_MAX_QUANTITY_WHOLE = "1000000"


class ImportFileError(ValueError):
    """An upload or workbook error safe to show in the import form."""


@dataclass
class PortfolioRow:
    asset_code: str = ""
    asset_name: str = ""
    quantity: Decimal | None = None
    cost_cents: int | None = None
    market_value_cents: int | None = None
    available_cash_cents: int | None = None
    occurred_on: date | None = None
    source_text: str = ""
    confidence: float | None = None
    errors: list[str] = field(default_factory=list)
    editable_values: dict[str, str] = field(default_factory=dict)

    def form_values(self) -> dict[str, str]:
        return self.editable_values or {
            "asset_code": self.asset_code,
            "asset_name": self.asset_name,
            "quantity": "" if self.quantity is None else str(self.quantity),
            "cost": _format_cents(self.cost_cents),
            "market_value": _format_cents(self.market_value_cents),
            "available_cash": _format_cents(self.available_cash_cents),
            "date": "" if self.occurred_on is None else self.occurred_on.isoformat(),
        }


def _format_cents(value: int | None) -> str:
    if value is None:
        return ""
    return f"{value // 100}.{value % 100:02d}"


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _money(value: str, label: str, errors: list[str]) -> int | None:
    if value == "":
        return None
    if _MONEY.fullmatch(value) is None:
        errors.append(f"{label}必须是最多两位小数的非负金额")
        return None
    whole, dot, fraction = value.partition(".")
    cents_text = (whole + ((fraction + "00")[:2] if dot else "00")).lstrip("0")
    cents_text = cents_text or "0"
    if len(cents_text) > len(_MAX_CENTS_TEXT) or (
        len(cents_text) == len(_MAX_CENTS_TEXT) and cents_text > _MAX_CENTS_TEXT
    ):
        errors.append(f"{label}超出可记录范围")
        return None
    return int(cents_text)


def _quantity(value: str, errors: list[str]) -> Decimal | None:
    if value == "":
        return None
    if _QUANTITY.fullmatch(value) is None:
        errors.append("份额必须是最多八位小数的非负数字")
        return None
    whole, _, fraction = value.partition(".")
    if (
        len(whole) > len(_MAX_QUANTITY_WHOLE)
        or (
            len(whole) == len(_MAX_QUANTITY_WHOLE)
            and whole > _MAX_QUANTITY_WHOLE
        )
        or (whole == _MAX_QUANTITY_WHOLE and any(digit != "0" for digit in fraction))
    ):
        errors.append("份额超出可记录范围")
        return None
    try:
        parsed = Decimal(value)
    except InvalidOperation:
        errors.append("份额格式无效")
        return None
    return parsed


def _date(value: str, errors: list[str]) -> date | None:
    if value == "":
        return None
    try:
        return date.fromisoformat(value)
    except ValueError:
        errors.append("日期必须是 YYYY-MM-DD")
        return None


def normalize_row(
    values: Mapping[str, object],
    *,
    source_text: str = "",
    confidence: float | None = None,
) -> PortfolioRow:
    """Validate editable normalized fields into exact ledger-ready values."""
    text = {key: _cell_text(value) for key, value in values.items()}
    errors: list[str] = []
    code = text.get("asset_code", "")
    name = text.get("asset_name", "")
    quantity = _quantity(text.get("quantity", ""), errors)
    cost = _money(text.get("cost", ""), "成本金额", errors)
    market_value = _money(text.get("market_value", ""), "当前市值", errors)
    available_cash = _money(text.get("available_cash", ""), "可用现金", errors)
    occurred_on = _date(text.get("date", ""), errors)

    has_asset = any((code, name, text.get("quantity", ""), text.get("cost", "")))
    if has_asset:
        if _CODE.fullmatch(code) is None:
            errors.append("基金代码格式无效")
        if not name or len(name) > 200:
            errors.append("基金名称不能为空且不能超过 200 字符")
        if quantity is None or quantity <= 0:
            errors.append("份额必须大于 0")
        if cost is None or cost <= 0:
            errors.append("持仓成本必须大于 0")
    if available_cash is not None and available_cash < 0:
        errors.append("可用现金不能为负")
    if not has_asset and available_cash is None:
        errors.append("每行必须包含持仓或可用现金")
    if occurred_on is None:
        errors.append("日期不能为空")

    return PortfolioRow(
        asset_code=code,
        asset_name=name,
        quantity=quantity,
        cost_cents=cost,
        market_value_cents=market_value,
        available_cash_cents=available_cash,
        occurred_on=occurred_on,
        source_text=source_text,
        confidence=confidence,
        errors=list(dict.fromkeys(errors)),
        editable_values=text,
    )


def _normalized_values(raw: dict[str, object]) -> dict[str, object]:
    values: dict[str, object] = {}
    for heading, value in raw.items():
        normalized = _ALIASES.get(_cell_text(heading).replace(" ", ""))
        if normalized is not None:
            values[normalized] = value
    return values


def _row(raw: dict[str, object]) -> PortfolioRow:
    source_text = " | ".join(
        f"{_cell_text(key)}: {_cell_text(value)}" for key, value in raw.items()
    )
    return normalize_row(_normalized_values(raw), source_text=source_text)


def _bounded_texts(values: Sequence[object]) -> list[str]:
    texts = [_cell_text(value) for value in values]
    if any(len(text) > MAX_IMPORT_FIELD_CHARACTERS for text in texts):
        raise ImportFileError("import field limit exceeded")
    return texts


def _check_columns(count: int) -> None:
    if count > MAX_IMPORT_COLUMNS:
        raise ImportFileError("import column limit exceeded")


def _check_cells(count: int) -> None:
    if count > MAX_IMPORT_CELLS:
        raise ImportFileError("import cell limit exceeded")


def _row_preview_characters(headers: list[str], values: list[str]) -> int:
    pair_count = min(len(headers), len(values))
    if pair_count == 0:
        return 0
    return sum(
        len(headers[index]) + 2 + len(values[index]) for index in range(pair_count)
    ) + 3 * (pair_count - 1)


def _check_preview_characters(count: int) -> None:
    if count > MAX_PREVIEW_CHARACTERS:
        raise ImportFileError("import preview size limit exceeded")


def _csv_rows(content: bytes) -> list[dict[str, object]]:
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportFileError("CSV must be UTF-8 encoded") from exc
    try:
        reader = csv.reader(StringIO(decoded), strict=True)
        try:
            raw_headers = next(reader)
        except StopIteration as exc:
            raise ImportFileError("CSV is missing a header row") from exc
        headers = _bounded_texts(raw_headers)
        if not headers:
            raise ImportFileError("CSV is missing a header row")
        _check_columns(len(headers))
        cells = len(headers)
        _check_cells(cells)
        preview_characters = 0
        rows: list[dict[str, object]] = []
        for row_number, raw_values in enumerate(reader, start=1):
            if row_number > MAX_IMPORT_ROWS:
                raise ImportFileError("import row limit exceeded")
            if len(raw_values) > len(headers):
                raise ImportFileError("import column limit exceeded")
            _check_columns(len(raw_values))
            cells += len(raw_values)
            _check_cells(cells)
            values = _bounded_texts(raw_values)
            preview_characters += _row_preview_characters(headers, values)
            _check_preview_characters(preview_characters)
            if any(values):
                rows.append(dict(zip(headers, raw_values, strict=False)))
        return rows
    except csv.Error as exc:
        raise ImportFileError("CSV could not be read") from exc


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
    except (BadZipFile, OSError) as exc:
        raise ImportFileError("XLSX workbook could not be read") from exc

    if sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ImportFileError("XLSX expanded size limit exceeded")
    for entry in entries:
        if entry.file_size > MAX_XLSX_ENTRY_BYTES:
            raise ImportFileError("XLSX entry size limit exceeded")
        if (
            entry.file_size >= _MIN_RATIO_CHECK_BYTES
            and entry.file_size / max(entry.compress_size, 1)
            > MAX_XLSX_COMPRESSION_RATIO
        ):
            raise ImportFileError("XLSX compression ratio limit exceeded")


def _xlsx_rows(content: bytes) -> list[dict[str, object]]:
    _validate_xlsx_archive(content)
    workbook = None
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        declared_rows = sheet.max_row or 0
        declared_columns = sheet.max_column or 0
        if declared_rows > MAX_IMPORT_ROWS + 1:
            raise ImportFileError("import row limit exceeded")
        _check_columns(declared_columns)
        _check_cells(declared_rows * declared_columns)
        iterator = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(iterator)
        except StopIteration as exc:
            raise ImportFileError("XLSX workbook could not be read") from exc
        headers = _bounded_texts(raw_headers)
        _check_columns(len(headers))
        cells = len(headers)
        _check_cells(cells)
        preview_characters = 0
        rows: list[dict[str, object]] = []
        for row_number, raw_values in enumerate(iterator, start=1):
            if row_number > MAX_IMPORT_ROWS:
                raise ImportFileError("import row limit exceeded")
            if len(raw_values) > len(headers):
                raise ImportFileError("import column limit exceeded")
            _check_columns(len(raw_values))
            cells += len(raw_values)
            _check_cells(cells)
            values = _bounded_texts(raw_values)
            preview_characters += _row_preview_characters(headers, values)
            _check_preview_characters(preview_characters)
            if any(values):
                rows.append(dict(zip(headers, raw_values, strict=False)))
        return rows
    except ImportFileError:
        raise
    except (
        BadZipFile,
        OSError,
        ParseError,
        ValueError,
        KeyError,
        StopIteration,
    ) as exc:
        raise ImportFileError("XLSX workbook could not be read") from exc
    finally:
        if workbook is not None:
            workbook.close()


def parse_portfolio_file(content: bytes, filename: str) -> list[PortfolioRow]:
    """Parse an allowed structured upload; image files are handled by OCR."""
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ImportFileError("unsupported import file type")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ImportFileError("upload exceeds the 10 MiB limit")
    if extension not in {".csv", ".xlsx"}:
        raise ImportFileError("image uploads must be parsed by OCR")
    raw_rows = _csv_rows(content) if extension == ".csv" else _xlsx_rows(content)
    if not raw_rows:
        raise ImportFileError("import file contains no data rows")
    return [_row(raw) for raw in raw_rows]
